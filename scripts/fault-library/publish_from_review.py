#!/usr/bin/env python3
"""
Fault library: turn a COMPLETED technical review pack (XLSX) into a guarded publish SQL file.

Nothing is written to the database by this script. It:
  1. reads the reviewer's decisions from the "Review" sheet,
  2. reads the current draft rows from the database (psql, read-only),
  3. applies the approval rules below,
  4. writes a report CSV and a single-transaction SQL file for separate, explicit approval.

Rules
  - Publish ONLY rows with Correct? = "Y", reviewer name/RGI and date (DD/MM/YY) filled,
    Review state = "Ready for technical review", DB status still 'draft', and the
    explanation + causes in the pack identical to the database (no edits since review).
  - Correct? = "N" (with a comment) -> status 'rejected'.
  - Everything else (blank, unresolved, mismatched, missing sign-off) stays 'draft'.
  - A model is published only if at least one of its codes is published. Row-level
    security still hides every non-published code of a published model.
  - Every UPDATE is guarded (id + status='draft' + unchanged text), so re-running is safe.

Usage
  python3 scripts/fault-library/publish_from_review.py <completed.xlsx> <approver_auth_uuid> <out_dir>
"""
import csv, datetime, hashlib, json, os, re, subprocess, sys
from openpyxl import load_workbook

READY = "Ready for technical review"
DATE_RE = re.compile(r"^\d{2}/\d{2}/\d{2}$")
UUID_RE = re.compile(r"^[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}$")


def q(s):
    return "'" + str(s).replace("'", "''") + "'"


def fp(explanation, causes):
    return hashlib.sha256((explanation.strip() + "\x1f" + causes.strip()).encode()).hexdigest()


def db_rows():
    sql = ("select json_agg(json_build_object('id',c.id,'model_id',m.id,'model',m.model_name,'code',c.code,"
           "'status',c.status,'explanation',c.explanation,'causes',array_to_string(c.possible_causes,' | ')))"
           " from boiler_fault_codes c join boiler_fault_models m on m.id=c.model_id")
    out = subprocess.run(["psql", "-At", "-c", sql], capture_output=True, text=True, check=True).stdout.strip()
    return {(r["model"], r["code"]): r for r in json.loads(out or "[]")}


def main(path, approver, out_dir):
    if not UUID_RE.match(approver):
        sys.exit("approver must be the approving superadmin's auth user id (uuid)")
    ws = load_workbook(path, data_only=True)["Review"]
    head = [str(c.value or "").strip() for c in ws[1]]
    col = {h: i for i, h in enumerate(head)}
    need = ["Model", "Code / message", "Explanation", "Possible causes / engineer checks", "Review state",
            "Correct? (Y/N)", "Engineer comment", "Reviewer name / RGI no.", "Date (DD/MM/YY)"]
    missing = [h for h in need if h not in col]
    if missing:
        sys.exit(f"Review sheet missing columns: {missing}")

    db = db_rows()
    seen, report, pub, rej, models = set(), [], [], [], {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        g = lambda h: "" if row[col[h]] is None else str(row[col[h]]).strip()
        key = (g("Model"), g("Code / message"))
        if key in seen:
            sys.exit(f"Duplicate row in pack: {key}")
        seen.add(key)
        r = db.get(key)
        verdict, why = "keep draft", ""
        ans = g("Correct? (Y/N)").upper()
        if r is None:
            verdict, why = "ignored", "not a library entry (gap/placeholder row)"
        elif r["status"] != "draft":
            verdict, why = "ignored", f"already {r['status']}"
        elif fp(g("Explanation"), g("Possible causes / engineer checks")) != fp(r["explanation"], r["causes"]):
            verdict, why = "keep draft", "pack text differs from database - re-review needed"
        elif ans == "N":
            if not g("Engineer comment"):
                why = "rejected without comment - held as draft"
            else:
                verdict = "reject"
        elif ans == "Y":
            if g("Review state") != READY:
                why = "unresolved entry - needs separate explicit approval"
            elif not g("Reviewer name / RGI no.") or not DATE_RE.match(g("Date (DD/MM/YY)")):
                why = "missing reviewer name/RGI or DD/MM/YY date"
            else:
                verdict = "publish"
        elif ans:
            why = f"unrecognised answer '{ans}'"
        else:
            why = "not checked"
        report.append([key[0], key[1], ans, g("Reviewer name / RGI no."), g("Date (DD/MM/YY)"), verdict, why])
        if verdict == "publish":
            pub.append((r, g("Reviewer name / RGI no."), g("Date (DD/MM/YY)")))
            models[r["model_id"]] = key[0]
        elif verdict == "reject":
            rej.append((r, g("Engineer comment"), g("Reviewer name / RGI no."), g("Date (DD/MM/YY)")))

    os.makedirs(out_dir, exist_ok=True)
    with open(os.path.join(out_dir, "publish_report.csv"), "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Model", "Code", "Correct?", "Reviewer", "Date", "Decision", "Reason"])
        w.writerows(report)

    stamp = datetime.date.today().strftime("%d/%m/%y")
    lines = ["begin;"]
    guard = lambda r: (f"id={q(r['id'])} and status='draft' and explanation={q(r['explanation'])} "
                       f"and array_to_string(possible_causes,' | ')={q(r['causes'])}")
    for r, who, d in pub:
        note = f"\nSIGN-OFF: approved by {who} on {d}; published {stamp}."
        lines.append(f"update public.boiler_fault_codes set status='published', verified_by={q(approver)}, "
                     f"verified_at=now(), technical_details=coalesce(technical_details,'')||{q(note)} where {guard(r)};")
    for r, why, who, d in rej:
        note = f"\nREJECTED by {who} on {d}: {why}"
        lines.append(f"update public.boiler_fault_codes set status='rejected', "
                     f"technical_details=coalesce(technical_details,'')||{q(note)} where {guard(r)};")
    for mid in models:
        lines.append(f"update public.boiler_fault_models set status='published', verified_by={q(approver)}, verified_at=now() "
                     f"where id={q(mid)} and status='draft' and exists (select 1 from public.boiler_fault_codes "
                     f"where model_id={q(mid)} and status='published');")
    lines.append("commit;")
    with open(os.path.join(out_dir, "publish.sql"), "w") as f:
        f.write("\n".join(lines) + "\n")

    counts = {k: sum(1 for x in report if x[5] == k) for k in ("publish", "reject", "keep draft", "ignored")}
    print(json.dumps({"counts": counts, "models_to_publish": sorted(models.values())}, indent=1))


if __name__ == "__main__":
    if len(sys.argv) != 4:
        sys.exit(__doc__)
    main(*sys.argv[1:])
